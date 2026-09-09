import streamlit as st
import pandas as pd
import openpyxl
import datetime
import os
import re
import json

EXCLUSION_FILE = "excluded_faculty.json"

st.set_page_config(
    page_title="UID ID Timetable & Generator",
    page_icon="📅",
    layout="wide"
)

st.markdown("""
<style>
    .badge-free {
        background-color: #dcfce7;
        color: #15803d;
        padding: 4px 10px;
        border-radius: 6px;
        font-weight: 600;
        font-size: 0.85rem;
        display: inline-block;
        margin: 2px 0;
    }
    .badge-busy {
        background-color: #fee2e2;
        color: #b91c1c;
        padding: 4px 10px;
        border-radius: 6px;
        font-weight: 600;
        font-size: 0.85rem;
        display: inline-block;
        margin: 2px 0;
    }
    .faculty-card {
        background: #f8fafc;
        border: 1px solid #e2e8f0;
        border-radius: 8px;
        padding: 12px 16px;
        margin-bottom: 8px;
    }
    .wt-table {
        width: 100%;
        border-collapse: collapse;
        font-family: inherit;
        margin-top: 14px;
    }
    .wt-th {
        background-color: #0f172a;
        color: #ffffff;
        padding: 10px 8px;
        text-align: center;
        font-size: 0.88rem;
        border: 1px solid #334155;
    }
    .wt-td-label {
        background-color: #f8fafc;
        color: #0f172a;
        font-weight: 700;
        padding: 8px 10px;
        border: 1px solid #cbd5e1;
        font-size: 0.85rem;
        vertical-align: middle;
        white-space: nowrap;
    }
    .wt-td {
        border: 1px solid #cbd5e1;
        padding: 6px 4px;
        vertical-align: top;
        background-color: #ffffff;
        width: 18%;
    }
    .s-pill {
        border-radius: 4px;
        padding: 3px 6px;
        margin-bottom: 4px;
        font-size: 0.74rem;
        line-height: 1.25;
        display: block;
        border-left: 3px solid transparent;
    }
    .s1 { background: #dbeafe; color: #1e40af; border-color: #2563eb; }
    .s2 { background: #ccfbf1; color: #0f766e; border-color: #0d9488; }
    .s3 { background: #fef3c7; color: #92400e; border-color: #d97706; }
    .s4 { background: #f3e8ff; color: #6b21a8; border-color: #9333ea; }
    .s-off { background: #f1f5f9; color: #94a3b8; font-style: italic; }
</style>
""", unsafe_allow_html=True)


def get_excluded_faculty():
    if os.path.exists(EXCLUSION_FILE):
        try:
            with open(EXCLUSION_FILE, "r") as f:
                return json.load(f)
        except Exception:
            return []
    return []


COURSES_DATA = [
    # 1st Sem B.Design (PSD) Batch 2026-30
    {"code": "26UDU01001", "title": "Design Essentials", "sem": "I", "prog_code": "1019", "batch": "1st Sem B.Des (PSD) 2026-30"},
    {"code": "26UDU01002", "title": "Visual Representation Skills", "sem": "I", "prog_code": "1019", "batch": "1st Sem B.Des (PSD) 2026-30"},
    {"code": "26UDU01103", "title": "Materials and Craftsmanship I", "sem": "I", "prog_code": "1019", "batch": "1st Sem B.Des (PSD) 2026-30"},
    {"code": "26UDU01204", "title": "Colour and Form", "sem": "I", "prog_code": "1019", "batch": "1st Sem B.Des (PSD) 2026-30"},
    {"code": "26UDU01305", "title": "History of Art & Design", "sem": "I", "prog_code": "1019", "batch": "1st Sem B.Des (PSD) 2026-30"},
    {"code": "26UDU01406", "title": "Fundamentals of AI", "sem": "I", "prog_code": "1019", "batch": "1st Sem B.Des (PSD) 2026-30"},
    {"code": "26UDU01607", "title": "Digital Tools", "sem": "I", "prog_code": "1019", "batch": "1st Sem B.Des (PSD) 2026-30"},
    
    # 3rd Sem B. Design Product Design Batch 2025-29
    {"code": "UC012030001", "title": "Personality Development", "sem": "III", "prog_code": "1019", "batch": "3rd Sem B.Des PD 2025-29"},
    {"code": "31203001203", "title": "Design Research", "sem": "III", "prog_code": "1019", "batch": "3rd Sem B.Des PD 2025-29"},
    {"code": "31203006211", "title": "Product Visualization", "sem": "III", "prog_code": "1019", "batch": "3rd Sem B.Des PD 2025-29"},
    {"code": "31203006212", "title": "Form, Aesthetic and Emotion", "sem": "III", "prog_code": "1019", "batch": "3rd Sem B.Des PD 2025-29"},
    {"code": "31203006207", "title": "Studio- Human Centric Design", "sem": "III", "prog_code": "1019", "batch": "3rd Sem B.Des PD 2025-29"},
    {"code": "31203006215", "title": "Design Articulation with AI", "sem": "III", "prog_code": "1019", "batch": "3rd Sem B.Des PD 2025-29"},
    {"code": "31203006214", "title": "Indian Design System", "sem": "III", "prog_code": "1019", "batch": "3rd Sem B.Des PD 2025-29"},
    
    # 5th Sem B. Design Product Design Batch 2024-28
    {"code": "31305006328", "title": "CAID", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001303", "title": "Conceptualization & Characterization", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001304", "title": "Speed Modelling in Clay", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001306", "title": "Space Perception", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001324", "title": "Matte-Painting", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001325", "title": "Ad Film Production", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001326", "title": "Fashion Styling", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001327", "title": "UX/UI", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001328", "title": "The Art of Delightful Design", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305001329", "title": "Branding And Identity Design", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305006322", "title": "Human Factors", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305006329", "title": "Portfolio Design with Voice Agents", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305006324", "title": "Studio- Humanizing Technology", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305006325", "title": "Experience Design", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305006326", "title": "Packaging Design", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305006327", "title": "Speculative design", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    {"code": "31305006329", "title": "Portfolio with AI", "sem": "V", "prog_code": "1019", "batch": "5th Sem B.Des PD 2024-28"},
    
    # 7th Sem B.Design Product Design Batch 2023-27
    {"code": "31407001400", "title": "Design Management", "sem": "VII", "prog_code": "1019", "batch": "7th Sem B.Des PD 2023-27"},
    {"code": "31407006403", "title": "Internship", "sem": "VII", "prog_code": "1019", "batch": "7th Sem B.Des PD 2023-27"},
    {"code": "31407006405", "title": "System Analysis and Design with Voice Agents", "sem": "VII", "prog_code": "1019", "batch": "7th Sem B.Des PD 2023-27"},
    
    # 1st Sem M. Design Product & Service Design Batch 2026-28
    {"code": "26UDP01317", "title": "Design Appreciation & Storytelling", "sem": "I", "prog_code": "1025", "batch": "1st Sem M.Des PSD 2026-28"},
    {"code": "26UDP01018", "title": "Design Foundation", "sem": "I", "prog_code": "1025", "batch": "1st Sem M.Des PSD 2026-28"},
    {"code": "26UDP01019", "title": "Form Studies", "sem": "I", "prog_code": "1025", "batch": "1st Sem M.Des PSD 2026-28"},
    {"code": "26UDP01020", "title": "Design Studio I", "sem": "I", "prog_code": "1025", "batch": "1st Sem M.Des PSD 2026-28"},
    {"code": "26UDP01421", "title": "Design Prototyping", "sem": "I", "prog_code": "1025", "batch": "1st Sem M.Des PSD 2026-28"},
    {"code": "26UDP01122", "title": "Frugal Innovation", "sem": "I", "prog_code": "1025", "batch": "1st Sem M.Des PSD 2026-28"},
    {"code": "26UDP01123", "title": "Emergent Technology", "sem": "I", "prog_code": "1025", "batch": "1st Sem M.Des PSD 2026-28"},
    {"code": "26UDP01624", "title": "CAID", "sem": "I", "prog_code": "1025", "batch": "1st Sem M.Des PSD 2026-28"},
    
    # 3rd Sem Masters in Industrial Design Batch 2025-27
    {"code": "32203001602", "title": "Internship", "sem": "III", "prog_code": "1025", "batch": "3rd Sem M.Des ID 2025-27"},
    {"code": "32203001603", "title": "Entrepreneurship", "sem": "III", "prog_code": "1025", "batch": "3rd Sem M.Des ID 2025-27"},
    {"code": "32203001604", "title": "Research Methodology", "sem": "III", "prog_code": "1025", "batch": "3rd Sem M.Des ID 2025-27"},
    {"code": "32203002613", "title": "Generative AI for UI & UX Design", "sem": "III", "prog_code": "1025", "batch": "3rd Sem M.Des ID 2025-27"},
    {"code": "32203002610", "title": "Studio: Design and Technology", "sem": "III", "prog_code": "1025", "batch": "3rd Sem M.Des ID 2025-27"},
    {"code": "32203002611", "title": "Lighting Design", "sem": "III", "prog_code": "1025", "batch": "3rd Sem M.Des ID 2025-27"},
    {"code": "32203002612", "title": "Craft and Technology", "sem": "III", "prog_code": "1025", "batch": "3rd Sem M.Des ID 2025-27"},
]

DEFAULT_FACULTY_LIST = {
    "-- None / Leave Empty --": "",
    "✏️ [Manual / Custom Entry]": "CUSTOM",
    "Sharad Shekar Shetty (15044)": "15044",
    "Kishan Chavda (15177)": "15177",
    "Umang Shah (15069)": "15069",
    "Shoeb Iqbal Khan (15255)": "15255",
    "Abhishek Karmakar (15254)": "15254",
    "Malekulashter (15342)": "15342",
    "Arshkirat Singh Gill (15279)": "15279",
    "Aditya Lingam (15067)": "15067",
    "Sreya Acharyya (15461)": "15461",
    "Dhanush Kumar (15497)": "15497",
    "Vipul Vinayak Jadhav (15500)": "15500",
    "Ashuj Chawda (15506)": "15506",
    "DA Siddharth (15445)": "15445",
    "Sree Hari B Lal (15546)": "15546",
    "Mithil Suresh (15552)": "15552",
    "Bhargav Manchalla (15556)": "15556",
    "Atul Kedia (15477)": "15477",
    "Bhargav Mistry (15478)": "15478",
}

COURSE_OPTIONS = {f"{c['title']}  |  [{c['code']}]  ({c['batch']})": c for c in COURSES_DATA}
COURSE_OPTIONS["✏️ [Custom / Manual Course Entry]"] = {
    "code": "", "title": "Custom", "sem": "I", "prog_code": "1019", "batch": "Custom"
}
SEMESTERS = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
DAYS_MAP = {"Mon": 0, "Tue": 1, "Wed": 2, "Thu": 3, "Fri": 4, "Sat": 5}
SLOT_TIMINGS_MAP = {
    1: "09:30 TO 11:00",
    2: "11:20 TO 13:10",
    3: "14:05 TO 15:40",
    4: "16:00 TO 17:30"
}


@st.cache_data
def load_master_data():
    files = [f for f in os.listdir('.') if f.endswith('.xlsx') and not f.startswith('~')]
    if not files:
        st.error("No `.xlsx` timetable spreadsheet found in the directory.")
        st.stop()
    
    file_path = files[0]
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    df_fw = pd.read_excel(file_path, sheet_name='Faculty Work Load ')
    raw_faculties = df_fw['Faculty Name '].dropna().unique().tolist()
    
    official_faculty_list = []
    first_name_to_full = {}
    first_name_to_code = {}

    for f in raw_faculties:
        full = " ".join(str(f).split())
        tokens = [t.lower().strip("().,") for t in full.split()]
        tokens = [t for t in tokens if t not in ['dr', 'mr', 'ms', 'prof']]
        if tokens:
            fn = tokens[0]
            first_name_to_full[fn] = full
            official_faculty_list.append(full)
            for k_erp, code_erp in DEFAULT_FACULTY_LIST.items():
                if fn in k_erp.lower():
                    first_name_to_code[fn] = code_erp
                    break
            
    official_faculty_list = sorted(list(set(official_faculty_list)))
    
    sheet_name = 'Morning_Afternoon Updated_Timet' if 'Morning_Afternoon Updated_Timet' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    month_map = {'Jul': 7, 'Aug': 8, 'Sep': 9, 'Oct': 10, 'Nov': 11, 'Dec': 12}
    calendar_cols = []

    for col in range(2, ws.max_column + 1):
        m_val = ws.cell(3, col).value
        d_val = ws.cell(4, col).value
        w_val = ws.cell(5, col).value

        if d_val is not None and str(d_val).strip() not in ['', 'None', 'Date']:
            try:
                day_int = int(float(d_val))
                m_str = str(m_val).strip()
                if m_str in month_map:
                    w_str = str(w_val).strip() if w_val else "Week"
                    d_obj = datetime.date(2026, month_map[m_str], day_int)
                    calendar_cols.append({
                        'col': col,
                        'date': d_obj,
                        'week': w_str,
                        'label': f"{d_obj.strftime('%a, %d %b %Y')} ({w_str})"
                    })
            except Exception:
                continue

    teaching_rows = [
        # UG-Sem 3
        (14, 'UG-Sem 3', 'A', 'Morning', 8),
        (15, 'UG-Sem 3', 'A', 'Afternoon', 8),
        (16, 'UG-Sem 3', 'B', 'Morning', 8),
        (17, 'UG-Sem 3', 'B', 'Afternoon', 8),
        (18, 'UG-Sem 3', 'C', 'Morning', 8),
        (19, 'UG-Sem 3', 'C', 'Afternoon', 8),
        (20, 'UG-Sem 3', 'D', 'Morning', 8),
        (21, 'UG-Sem 3', 'D', 'Afternoon', 8),
        (22, 'UG-Sem 3', 'E', 'Morning', 8),
        (23, 'UG-Sem 3', 'E', 'Afternoon', 8),
        # UG-Sem 5
        (33, 'UG-Sem 5', 'A', 'Morning', 30),
        (34, 'UG-Sem 5', 'A', 'Afternoon', 30),
        (35, 'UG-Sem 5', 'B', 'Morning', 30),
        (36, 'UG-Sem 5', 'B', 'Afternoon', 30),
        (37, 'UG-Sem 5', 'C', 'Morning', 30),
        (38, 'UG-Sem 5', 'C', 'Afternoon', 30),
        (39, 'UG-Sem 5', 'D', 'Morning', 30),
        (40, 'UG-Sem 5', 'D', 'Afternoon', 30),
        # UG-Sem 7
        (51, 'UG-Sem 7', 'A', 'Lead', 48),
        (52, 'UG-Sem 7', 'A', 'Assisting', 48),
        (53, 'UG-Sem 7', 'B', 'Lead', 48),
        (54, 'UG-Sem 7', 'B', 'Assisting', 48),
        (55, 'UG-Sem 7', 'C', 'Lead', 48),
        (56, 'UG-Sem 7', 'C', 'Assisting', 48),
        # PG-Sem 1
        (65, 'PG-Sem 1', 'Combined', 'Full Day', 62),
        # PG-Sem 3
        (74, 'PG-Sem 3', 'Combined', 'Full Day', 71)
    ]

    faculty_day_schedule = {f: {} for f in official_faculty_list}
    cohort_day_schedule = {}

    for c_info in calendar_cols:
        col = c_info['col']
        dt = c_info['date']
        if dt not in cohort_day_schedule:
            cohort_day_schedule[dt] = {}

        for r, cohort, sec, slot, mod_r in teaching_rows:
            cell_val = ws.cell(r, col).value
            if cell_val is not None and str(cell_val).strip() not in ['', 'None', '-']:
                txt = str(cell_val).strip()
                mod_name = str(ws.cell(mod_r, col).value or "").strip()
                clean_mod = " ".join(mod_name.split())
                
                cell_tokens = [re.sub(r'[^a-zA-Z]', '', w).lower() for w in txt.split()]
                matched_faculty = None
                faculty_code = None

                for tok in cell_tokens:
                    if tok in first_name_to_full:
                        matched_faculty = first_name_to_full[tok]
                        faculty_code = first_name_to_code.get(tok, "")
                        break

                if not matched_faculty:
                    matched_faculty = txt
                    faculty_code = "VISITING"

                if matched_faculty not in faculty_day_schedule:
                    faculty_day_schedule[matched_faculty] = {}

                detail = f"{cohort} | Sec {sec} ({slot}) - {clean_mod}"
                if dt not in faculty_day_schedule[matched_faculty]:
                    faculty_day_schedule[matched_faculty][dt] = []
                if detail not in faculty_day_schedule[matched_faculty][dt]:
                    faculty_day_schedule[matched_faculty][dt].append(detail)

                if cohort not in cohort_day_schedule[dt]:
                    cohort_day_schedule[dt][cohort] = {}
                if sec not in cohort_day_schedule[dt][cohort]:
                    cohort_day_schedule[dt][cohort][sec] = {
                        "module": clean_mod,
                        "morning": None,
                        "afternoon": None,
                        "faculty_name_m": None,
                        "faculty_name_a": None
                    }

                if slot in ['Morning', 'Full Day', 'Lead']:
                    cohort_day_schedule[dt][cohort][sec]["morning"] = faculty_code
                    cohort_day_schedule[dt][cohort][sec]["faculty_name_m"] = matched_faculty
                if slot in ['Afternoon', 'Full Day', 'Assisting']:
                    cohort_day_schedule[dt][cohort][sec]["afternoon"] = faculty_code
                    cohort_day_schedule[dt][cohort][sec]["faculty_name_a"] = matched_faculty

    return calendar_cols, official_faculty_list, faculty_day_schedule, cohort_day_schedule


try:
    calendar_cols, official_faculty_list, faculty_schedule, cohort_day_schedule = load_master_data()
except Exception as e:
    st.error(f"Error loading master dataset: {e}")
    st.stop()

st.title("UID Department of Industrial Design")
st.caption("Master Academic Schedule, Faculty Availability & Timetable Generator")

tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "🔍 Faculty Availability", 
    "👤 Individual Schedule",
    "🗓️ Weekly Timetable",
    "⚡ Auto Timetable Export",
    "🛠️ Custom Generator Form"
])


# ==========================================================
# TAB 1: AVAILABILITY ENGINE
# ==========================================================
with tab1:
    excluded_faculties = set(get_excluded_faculty())

    date_to_week = {c['date']: c['week'] for c in calendar_cols}
    valid_dates = [c['date'] for c in calendar_cols]
    min_date = valid_dates[0]
    max_date = valid_dates[-1]

    default_date = datetime.date(2026, 9, 7) if min_date <= datetime.date(2026, 9, 7) <= max_date else min_date

    selected_date = st.date_input(
        "Select Academic Date:",
        value=default_date,
        min_value=min_date,
        max_value=max_date,
        key="avail_date_picker"
    )

    current_week = date_to_week.get(selected_date, "Non-Instructional / Holiday")

    busy_members = []
    busy_names = set()
    for fac, dates in faculty_schedule.items():
        classes = dates.get(selected_date, [])
        if classes and not any(skip in fac.lower() for skip in ['mid term', 'no faculty', 'workshop', 'tours']):
            busy_members.append((fac, classes))
            busy_names.add(fac)

    free_members = []
    for fac in official_faculty_list:
        if fac not in busy_names and fac not in excluded_faculties:
            free_members.append(fac)

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Week", current_week)
    m2.metric("Date", selected_date.strftime('%d %B %Y'))
    m3.metric("Available / Free", len(free_members))
    m4.metric("Scheduled in Class", len(busy_members))

    st.markdown("---")
    col_free, col_busy = st.columns([1, 1.2])

    with col_free:
        st.subheader(f"✅ Free Faculty ({len(free_members)})")
        st.caption("Official department faculty available for substitution:")
        f1, f2 = st.columns(2)
        for idx, fac in enumerate(free_members):
            target = f1 if idx % 2 == 0 else f2
            target.markdown(f"• **{fac}** &nbsp; <span class='badge-free'>Available</span>", unsafe_allow_html=True)

    with col_busy:
        st.subheader(f"🔒 Scheduled Faculty ({len(busy_members)})")
        st.caption("Instructors with active class assignments:")
        for fac, sessions in busy_members:
            details_html = "<br>".join([f"&nbsp;&nbsp;↳ {s}" for s in sessions])
            st.markdown(f"""
            <div class="faculty-card">
                <b>{fac}</b> &nbsp; <span class="badge-busy">In Class</span><br>
                <small style="color: #475569;">{details_html}</small>
            </div>
            """, unsafe_allow_html=True)


# ==========================================================
# TAB 2: INDIVIDUAL LOOKUP
# ==========================================================
with tab2:
    selected_fac = st.selectbox("Select Faculty Member:", options=official_faculty_list, key="indiv_fac_select")
    sched = faculty_schedule.get(selected_fac, {})

    if sched:
        st.markdown(f"**Total Teaching Days:** `{len(sched)} days`")
        records = []
        for c in calendar_cols:
            if c['date'] in sched:
                records.append({
                    'Date': c['date'].strftime('%d-%b-%Y'),
                    'Week': c['week'],
                    'Assigned Sessions': "; ".join(sched[c['date']])
                })
        st.table(pd.DataFrame(records))
    else:
        st.info(f"No active teaching assignments found for {selected_fac}.")


# ==========================================================
# TAB 3: WEEKLY CALENDAR TIMETABLE VIEW
# ==========================================================
with tab3:
    st.subheader("🗓️ Weekly Department Timetable")
    st.caption("Comprehensive weekly grid displaying all sections with Sessions 1–4 color-coded.")

    weeks_dict = {}
    for c_info in calendar_cols:
        w = c_info['week']
        if w not in weeks_dict:
            weeks_dict[w] = []
        weeks_dict[w].append(c_info)

    week_options = []
    week_data_map = {}
    for w, days in weeks_dict.items():
        d_start = days[0]['date'].strftime('%d %b')
        d_end = days[-1]['date'].strftime('%d %b %Y')
        lbl = f"{w}  ({d_start} - {d_end})"
        week_options.append(lbl)
        week_data_map[lbl] = (w, days)

    def_idx = next((i for i, opt in enumerate(week_options) if "WEEK 11" in opt), 0)

    if "selected_cohort" not in st.session_state:
        st.session_state["selected_cohort"] = "UG-Sem 3"

    col_week, col_cohorts = st.columns([1.1, 2.4])

    with col_week:
        selected_week_lbl = st.selectbox(
            "Select Week:",
            options=week_options,
            index=def_idx,
            key="weekly_tab_select"
        )
        selected_w_name, week_days = week_data_map[selected_week_lbl]

    with col_cohorts:
        st.markdown("<p style='font-size: 0.88rem; font-weight: 500; margin-bottom: 6px;'>Select Cohort:</p>", unsafe_allow_html=True)
        cohort_buttons = ["UG-Sem 3", "UG-Sem 5", "UG-Sem 7", "PG-Sem 1", "PG-Sem 3", "All"]
        btn_cols = st.columns(len(cohort_buttons))

        for b_col, opt in zip(btn_cols, cohort_buttons):
            is_active = (st.session_state["selected_cohort"] == opt)
            if b_col.button(
                opt,
                type="primary" if is_active else "secondary",
                use_container_width=True,
                key=f"btn_ch_{opt}"
            ):
                st.session_state["selected_cohort"] = opt
                st.rerun()

    cohort_filter = "All Cohorts" if st.session_state["selected_cohort"] == "All" else st.session_state["selected_cohort"]

    display_days = [d for d in week_days if d['date'].weekday() < 5]

    cohort_sections = {
        "UG-Sem 3": ["A", "B", "C", "D", "E"],
        "UG-Sem 5": ["A", "B", "C", "D"],
        "UG-Sem 7": ["A", "B", "C"],
        "PG-Sem 1": ["Combined"],
        "PG-Sem 3": ["Combined"]
    }

    active_cohorts = [cohort_filter] if cohort_filter != "All Cohorts" else list(cohort_sections.keys())

    html = ['<table class="wt-table">']
    html.append('<thead><tr>')
    html.append('<th class="wt-th" style="width: 12%;">Cohort / Section</th>')
    for d in display_days:
        d_name = d['date'].strftime('%a')
        d_num = d['date'].strftime('%d %b')
        html.append(f'<th class="wt-th">{d_name}<br><small>{d_num}</small></th>')
    html.append('</tr></thead><tbody>')

    for ch in active_cohorts:
        for sec in cohort_sections.get(ch, []):
            row_lbl = f"{ch}<br><span style='color: #0284c7;'>Sec {sec}</span>"
            html.append('<tr>')
            html.append(f'<td class="wt-td-label">{row_lbl}</td>')

            for d in display_days:
                dt = d['date']
                day_ch_data = cohort_day_schedule.get(dt, {}).get(ch, {})
                sec_info = day_ch_data.get(sec, None)

                html.append('<td class="wt-td">')

                if sec_info:
                    mod_name = sec_info.get("module", "")
                    mod_short = (mod_name[:24] + '...') if len(mod_name) > 24 else mod_name
                    fac_m = sec_info.get("faculty_name_m") or "—"
                    fac_a = sec_info.get("faculty_name_a") or "—"
                    
                    is_thursday_off = (dt.weekday() == 3 and not ch.startswith("PG"))

                    html.append(f'<div class="s-pill s1"><b>S1:</b> {fac_m}<br><small>{mod_short}</small></div>')
                    html.append(f'<div class="s-pill s2"><b>S2:</b> {fac_m}<br><small>{mod_short}</small></div>')

                    if is_thursday_off:
                        html.append('<div class="s-pill s-off">S3 & S4: Off</div>')
                    else:
                        html.append(f'<div class="s-pill s3"><b>S3:</b> {fac_a}<br><small>{mod_short}</small></div>')
                        html.append(f'<div class="s-pill s4"><b>S4:</b> {fac_a}<br><small>{mod_short}</small></div>')
                else:
                    html.append('<div class="s-pill s-off" style="text-align: center; padding: 22px 0;">No Sessions</div>')

                html.append('</td>')
            html.append('</tr>')

    html.append('</tbody></table>')
    st.markdown("".join(html), unsafe_allow_html=True)


# ==========================================================
# TAB 4: WORKFLOW A — AUTO TIMETABLE EXPORT FROM EXCEL
# ==========================================================
with tab4:
    st.subheader("⚡ Auto-Generate Timetable CSV from Master Sheet")
    st.caption("Inspects the master spreadsheet, extracts the exact module and faculty assigned to each section, and maps it directly into ERP CSV format.")

    a_col1, a_col2, a_col3 = st.columns(3)
    with a_col1:
        cohort_choice = st.selectbox("Select Cohort:", ["UG-Sem 3", "UG-Sem 5", "UG-Sem 7", "PG-Sem 1", "PG-Sem 3"])
    with a_col2:
        auto_start = st.date_input("Start Date", datetime.date(2026, 9, 7), key="auto_start")
    with a_col3:
        auto_end = st.date_input("End Date", datetime.date(2026, 9, 11), key="auto_end")

    auto_thu_half = st.checkbox("Thursday Afternoon Off (Slots 3 & 4 off for UG)", value=True, key="auto_thu")
    default_block = st.text_input("Academic Block", value="F", key="auto_block")

    cohort_erp_defaults = {
        "UG-Sem 3": {"prog": "1019", "sem": "III"},
        "UG-Sem 5": {"prog": "1019", "sem": "V"},
        "UG-Sem 7": {"prog": "1019", "sem": "VII"},
        "PG-Sem 1": {"prog": "1025", "sem": "I"},
        "PG-Sem 3": {"prog": "1025", "sem": "III"},
    }

    if st.button("Auto-Extract & Build Timetable", type="primary", key="btn_auto_build"):
        if auto_start > auto_end:
            st.error("Start Date must be before or equal to End Date.")
        else:
            extracted_rows = []
            cur_date = auto_start

            while cur_date <= auto_end:
                if cur_date.weekday() < 5:
                    date_str = cur_date.strftime("%Y-%m-%d")
                    day_cohorts = cohort_day_schedule.get(cur_date, {})
                    cohort_data = day_cohorts.get(cohort_choice, {})
                    
                    is_thursday = (cur_date.weekday() == 3)
                    is_thu_off = is_thursday and auto_thu_half and not cohort_choice.startswith("PG")
                    active_slots = [1, 2] if is_thu_off else [1, 2, 3, 4]

                    for sec_code, sec_info in cohort_data.items():
                        mod_title = sec_info["module"]
                        fac_morning = sec_info["morning"]
                        fac_afternoon = sec_info["afternoon"]

                        matched_course = next((c for c in COURSES_DATA if c["title"].lower() in mod_title.lower() or mod_title.lower() in c["title"].lower()), None)
                        c_code = matched_course["code"] if matched_course else "3120300"
                        c_type = "MANDATORY"
                        c_classif = "PRACTICAL"

                        erp_info = cohort_erp_defaults.get(cohort_choice, {"prog": "1019", "sem": "III"})

                        for s_num in active_slots:
                            assigned_faculty = fac_morning if s_num in [1, 2] else fac_afternoon
                            extracted_rows.append({
                                "Date": date_str,
                                "Program Code": erp_info["prog"],
                                "Semester Code": erp_info["sem"],
                                "Year": None,
                                "Course Classification": c_classif,
                                "Course Code": c_code,
                                "Course Type": c_type,
                                "Faculty Code": assigned_faculty,
                                "Shift Timing": "09:30 TO 17:30",
                                "Slot Time": SLOT_TIMINGS_MAP[s_num],
                                "Section Code": sec_code,
                                "Group": None,
                                "Academic Block": default_block,
                                "Room Allocation": None,
                                "Combined Class": None,
                                "Mode of Class": "OFFLINE",
                                "Time Table Type": None
                            })

                cur_date += datetime.timedelta(days=1)

            if extracted_rows:
                df_auto = pd.DataFrame(extracted_rows)
                st.success(f"Extracted {len(df_auto)} rows directly from the spreadsheet!")
                st.table(df_auto)

                csv_auto = df_auto.to_csv(index=False).encode('utf-8')
                st.download_button(
                    label="📥 Download Auto-Generated CSV",
                    data=csv_auto,
                    file_name=f"Auto_Timetable_{cohort_choice}_{auto_start}.csv",
                    mime="text/csv"
                )
            else:
                st.warning(f"No scheduled sessions found for {cohort_choice} between {auto_start} and {auto_end}.")


# ==========================================================
# TAB 5: WORKFLOW B — MANUAL / FORM-BASED GENERATOR
# ==========================================================
with tab5:
    st.subheader("🛠️ Custom / Manual Timetable Generator")
    st.caption("Build and customize a timetable schedule using manual overrides, custom rooms, and configurable slots.")

    def render_course_form(label_prefix, key_suffix, default_index=0, default_classif="THEORY"):
        st.markdown(f"**{label_prefix}**")
        selected_key = st.selectbox(
            f"Subject for {label_prefix}",
            options=list(COURSE_OPTIONS.keys()),
            index=default_index,
            key=f"course_select_{key_suffix}"
        )
        course_info = COURSE_OPTIONS[selected_key]
        
        if selected_key == "✏️ [Custom / Manual Course Entry]":
            c1, c2 = st.columns(2)
            code = c1.text_input("Course Code", key=f"custom_code_{key_suffix}")
            ctype = c2.selectbox("Course Type", ["MANDATORY", "ELECTIVE"], index=0, key=f"ctype_{key_suffix}")
            c3, c4, c5 = st.columns(3)
            sem = c3.selectbox("Semester", SEMESTERS, index=0, key=f"sem_{key_suffix}")
            prog = c4.text_input("Program Code", value="1019", key=f"prog_{key_suffix}")
            classif = c5.selectbox("Classification", ["THEORY", "PRACTICAL"], index=0 if default_classif == "THEORY" else 1, key=f"classif_{key_suffix}")
        else:
            code = course_info["code"]
            default_sem_idx = SEMESTERS.index(course_info["sem"]) if course_info["sem"] in SEMESTERS else 0
            c1, c2, c3, c4 = st.columns([1.5, 1, 1, 1.2])
            ctype = c1.selectbox("Course Type", ["MANDATORY", "ELECTIVE"], index=0, key=f"ctype_{key_suffix}")
            sem = c2.selectbox("Semester", SEMESTERS, index=default_sem_idx, key=f"sem_{key_suffix}")
            prog = c3.text_input("Prog Code", value=course_info["prog_code"], key=f"prog_{key_suffix}")
            classif = c4.selectbox("Classification", ["THEORY", "PRACTICAL"], index=0 if default_classif == "THEORY" else 1, key=f"classif_{key_suffix}")
            st.caption(f"📌 Mapped: **Code:** `{code}` | **Sem:** `{sem}` | **Prog:** `{prog}`")
            
        return {"code": code, "type": ctype, "sem": sem, "prog": prog, "classification": classif}

    def render_faculty_dropdown(label, key_name):
        fac_label = st.selectbox(label, options=list(DEFAULT_FACULTY_LIST.keys()), key=f"fac_select_{key_name}")
        if fac_label == "✏️ [Manual / Custom Entry]":
            fac_code = st.text_input(f"Type Faculty Code ({label})", key=f"custom_fac_{key_name}")
        elif fac_label == "-- None / Leave Empty --":
            fac_code = ""
        else:
            fac_code = DEFAULT_FACULTY_LIST[fac_label]
        return fac_code.strip() if fac_code else None

    # Step 1: Schedule & Slots
    st.markdown("#### 1. Schedule & Slot Timings")
    t_col1, t_col2, t_col3 = st.columns([1.5, 1.5, 1])

    with t_col1:
        today = datetime.date.today()
        c_d1, c_d2 = st.columns(2)
        start_date = c_d1.date_input("Start Date", today, key="man_start")
        end_date = c_d2.date_input("End Date", today + datetime.timedelta(days=4), key="man_end")
        thursday_half_day = st.checkbox("Thursday Afternoon Off (Half Day)", value=True, key="man_thu")

    with t_col2:
        selected_days = st.multiselect("Active Days of Week", options=list(DAYS_MAP.keys()), default=["Mon", "Tue", "Wed", "Thu", "Fri"], key="man_days")

    with t_col3:
        academic_block = st.text_input("Academic Block", value="F", key="man_block")

    st.markdown("---")

    # Step 2: Course Configuration
    st.markdown("#### 2. Course Selection")
    c_top1, c_top2 = st.columns([2, 1])
    with c_top1:
        sync_morning = st.checkbox("🔄 Auto-fill Morning Slot 2 from Slot 1", value=True, key="man_sync_m")
        sync_afternoon = st.checkbox("🔄 Auto-fill Afternoon Slot 4 from Slot 3", value=True, key="man_sync_a")
    with c_top2:
        num_sections = st.slider("Number of Sections in Batch", min_value=1, max_value=5, value=4, key="man_sec_count")

    m_col, a_col = st.columns(2, gap="large")
    with m_col:
        st.markdown("### 🌅 Morning Slots")
        slot1_cfg = render_course_form("Slot 1 (09:30 - 11:00)", "man_s1", default_index=0, default_classif="THEORY")
        if sync_morning:
            s2_classif = st.selectbox("Slot 2 Classification", ["PRACTICAL", "THEORY"], index=0, key="man_s2_sync")
            slot2_cfg = {**slot1_cfg, "classification": s2_classif}
        else:
            slot2_cfg = render_course_form("Slot 2 (11:20 - 13:10)", "man_s2", default_index=0, default_classif="PRACTICAL")

    with a_col:
        st.markdown("### 🌇 Afternoon Slots")
        slot3_cfg = render_course_form("Slot 3 (14:05 - 15:40)", "man_s3", default_index=1, default_classif="THEORY")
        if sync_afternoon:
            s4_classif = st.selectbox("Slot 4 Classification", ["PRACTICAL", "THEORY"], index=0, key="man_s4_sync")
            slot4_cfg = {**slot3_cfg, "classification": s4_classif}
        else:
            slot4_cfg = render_course_form("Slot 4 (16:00 - 17:30)", "man_s4", default_index=1, default_classif="PRACTICAL")

    slots_configuration = {1: slot1_cfg, 2: slot2_cfg, 3: slot3_cfg, 4: slot4_cfg}

    st.markdown("---")

    # Step 3: Section Allocations
    st.markdown("#### 3. Section Allocations")
    section_labels = ["A", "B", "C", "D", "E"][:num_sections]
    section_allocations = []

    for sec in section_labels:
        st.markdown(f"**Section {sec}**")
        if sync_morning and sync_afternoon:
            col_m, col_a, col_r = st.columns([2.5, 2.5, 1.5])
            with col_m:
                fac_m = render_faculty_dropdown(f"Morning Faculty (Slots 1 & 2)", f"man_m_{sec}")
                fac_s1, fac_s2 = fac_m, fac_m
            with col_a:
                fac_a = render_faculty_dropdown(f"Afternoon Faculty (Slots 3 & 4)", f"man_a_{sec}")
                fac_s3, fac_s4 = fac_a, fac_a
            with col_r:
                room = st.text_input("Studio Room", key=f"man_room_{sec}", placeholder="e.g. F2, E10")
        else:
            col1, col2, col3, col4, col5 = st.columns([2, 2, 2, 2, 1.5])
            with col1:
                fac_s1 = render_faculty_dropdown("Slot 1 Fac", f"man_s1_{sec}")
            with col2:
                fac_s2 = fac_s1 if sync_morning else render_faculty_dropdown("Slot 2 Fac", f"man_s2_{sec}")
            with col3:
                fac_s3 = render_faculty_dropdown("Slot 3 Fac", f"man_s3_{sec}")
            with col4:
                fac_s4 = fac_s3 if sync_afternoon else render_faculty_dropdown("Slot 4 Fac", f"man_s4_{sec}")
            with col5:
                room = st.text_input("Studio Room", key=f"man_room_{sec}", placeholder="e.g. F2, E10")

        section_allocations.append({
            "section": sec,
            "fac_slot_1": fac_s1,
            "fac_slot_2": fac_s2,
            "fac_slot_3": fac_s3,
            "fac_slot_4": fac_s4,
            "room": room.strip() if room else None
        })

    st.markdown("---")

    # Step 4: Generation
    if st.button("Generate Manual Timetable", type="primary", use_container_width=True, key="man_gen_btn"):
        empty_courses = [f"Slot {k}" for k, v in slots_configuration.items() if not v["code"]]
        if empty_courses:
            st.error(f"Missing Course Code for: {', '.join(empty_courses)}")
        elif start_date > end_date:
            st.error("Start Date must be before or equal to End Date.")
        else:
            active_day_ints = [DAYS_MAP[d] for d in selected_days]
            rows = []
            cur_date = start_date
            
            while cur_date <= end_date:
                if cur_date.weekday() in active_day_ints:
                    is_thursday = (cur_date.weekday() == 3)
                    date_str = cur_date.strftime("%Y-%m-%d")
                    shift_timing = "09:30 TO 17:30"
                    active_slots = [1, 2] if (is_thursday and thursday_half_day) else [1, 2, 3, 4]
                    
                    for slot_num in active_slots:
                        s_cfg = slots_configuration[slot_num]
                        for sec in section_allocations:
                            rows.append({
                                "Date": date_str,
                                "Program Code": s_cfg["prog"],
                                "Semester Code": s_cfg["sem"],
                                "Year": None,
                                "Course Classification": s_cfg["classification"],
                                "Course Code": s_cfg["code"],
                                "Course Type": s_cfg["type"],
                                "Faculty Code": sec[f"fac_slot_{slot_num}"],
                                "Shift Timing": shift_timing,
                                "Slot Time": SLOT_TIMINGS_MAP[slot_num],
                                "Section Code": sec["section"],
                                "Group": None,
                                "Academic Block": academic_block,
                                "Room Allocation": sec["room"],
                                "Combined Class": None,
                                "Mode of Class": "OFFLINE",
                                "Time Table Type": None
                            })
                cur_date += datetime.timedelta(days=1)
            
            df_result = pd.DataFrame(rows)
            st.success(f"Generated {len(df_result)} schedule rows successfully!")
            st.table(df_result)
            
            csv_bytes = df_result.to_csv(index=False).encode('utf-8')
            sample_prog = slot1_cfg["prog"]
            sample_sem = slot1_cfg["sem"]
            
            st.download_button(
                label="📥 Download Timetable CSV (.csv)",
                data=csv_bytes,
                file_name=f"UID_Timetable_{sample_prog}_Sem{sample_sem}_{start_date}.csv",
                mime="text/csv",
                key="man_download_btn"
            )
